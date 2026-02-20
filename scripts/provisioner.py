#!/usr/bin/env python3
"""
Provisioner for ThingsBoard PE Industrial Historian.
Creates asset/device profiles, hierarchy, devices, relations, and attributes.

Usage:
    python provisioner.py --demo                    # Built-in demo data
    python provisioner.py --file sample_plant.xlsx  # From Excel
    python provisioner.py --generate-sample         # Generate sample xlsx
"""

import argparse
import json
import logging
import os
import sys
import time
from typing import Any, Dict, List, Optional, Tuple

import requests

# ---------------------------------------------------------------------------
# Logging setup
# ---------------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger("provisioner")

# ---------------------------------------------------------------------------
# Rate-limit helper
# ---------------------------------------------------------------------------
API_DELAY = 0.1  # seconds between API calls


def _sleep():
    """Rate-limit pause between API calls."""
    time.sleep(API_DELAY)


# ===================================================================
# ThingsBoard REST Client
# ===================================================================
class TBClient:
    """Minimal ThingsBoard PE REST client with idempotent helpers."""

    def __init__(self, base_url: str, username: str, password: str, dry_run: bool = False):
        self.base_url = base_url.rstrip("/")
        self.api = f"{self.base_url}/api"
        self.username = username
        self.password = password
        self.dry_run = dry_run
        self.token: Optional[str] = None
        self.session = requests.Session()

    # ------------------------------------------------------------------
    # Auth
    # ------------------------------------------------------------------
    def authenticate(self) -> None:
        """Obtain JWT token from ThingsBoard."""
        if self.dry_run:
            log.info("[DRY-RUN] Would authenticate as %s", self.username)
            self.token = "DRY_RUN_TOKEN"
            return

        url = f"{self.api}/auth/login"
        payload = {"username": self.username, "password": self.password}
        resp = requests.post(url, json=payload, timeout=30)
        resp.raise_for_status()
        self.token = resp.json()["token"]
        self.session.headers.update({
            "X-Authorization": f"Bearer {self.token}",
            "Content-Type": "application/json",
            "Accept": "application/json",
        })
        log.info("Authenticated as %s on %s", self.username, self.base_url)

    # ------------------------------------------------------------------
    # Low-level helpers
    # ------------------------------------------------------------------
    def _get(self, path: str, params: Optional[dict] = None) -> Any:
        _sleep()
        url = f"{self.api}{path}"
        resp = self.session.get(url, params=params, timeout=30)
        resp.raise_for_status()
        return resp.json()

    def _post(self, path: str, payload: Any = None) -> Any:
        _sleep()
        url = f"{self.api}{path}"
        resp = self.session.post(url, json=payload, timeout=30)
        resp.raise_for_status()
        try:
            return resp.json()
        except ValueError:
            return None

    # ------------------------------------------------------------------
    # Asset Profiles
    # ------------------------------------------------------------------
    def find_asset_profile(self, name: str) -> Optional[Dict]:
        """Search for an asset profile by name (paginated)."""
        if self.dry_run:
            return None
        page = 0
        while True:
            data = self._get("/assetProfiles", params={
                "pageSize": 100,
                "page": page,
                "sortProperty": "name",
                "sortOrder": "ASC",
            })
            for item in data.get("data", []):
                if item.get("name") == name:
                    return item
            if not data.get("hasNext", False):
                break
            page += 1
        return None

    def create_asset_profile(self, name: str, description: str = "") -> Dict:
        """Create or return existing asset profile."""
        existing = self.find_asset_profile(name)
        if existing:
            log.info("  Asset Profile '%s' already exists (id=%s)", name, existing["id"]["id"])
            return existing

        if self.dry_run:
            log.info("[DRY-RUN] Would create Asset Profile: %s", name)
            return {"id": {"id": f"dry-run-ap-{name}", "entityType": "ASSET_PROFILE"}, "name": name}

        payload = {
            "name": name,
            "description": description,
        }
        result = self._post("/assetProfile", payload)
        log.info("  Created Asset Profile '%s' (id=%s)", name, result["id"]["id"])
        return result

    # ------------------------------------------------------------------
    # Device Profiles
    # ------------------------------------------------------------------
    def find_device_profile(self, name: str) -> Optional[Dict]:
        """Search for a device profile by name (paginated)."""
        if self.dry_run:
            return None
        page = 0
        while True:
            data = self._get("/deviceProfiles", params={
                "pageSize": 100,
                "page": page,
                "sortProperty": "name",
                "sortOrder": "ASC",
            })
            for item in data.get("data", []):
                if item.get("name") == name:
                    return item
            if not data.get("hasNext", False):
                break
            page += 1
        return None

    def create_device_profile(self, name: str, description: str = "",
                              alarms: Optional[List] = None) -> Dict:
        """Create or return existing device profile with alarm rules."""
        existing = self.find_device_profile(name)
        if existing:
            log.info("  Device Profile '%s' already exists (id=%s)", name, existing["id"]["id"])
            return existing

        if self.dry_run:
            log.info("[DRY-RUN] Would create Device Profile: %s", name)
            return {"id": {"id": f"dry-run-dp-{name}", "entityType": "DEVICE_PROFILE"}, "name": name}

        payload = {
            "name": name,
            "description": description,
            "type": "DEFAULT",
            "transportType": "DEFAULT",
            "profileData": {
                "configuration": {"type": "DEFAULT"},
                "transportConfiguration": {"type": "DEFAULT"},
                "alarms": alarms or [],
                "provisionConfiguration": {
                    "type": "DISABLED",
                    "provisionDeviceSecret": None,
                },
            },
        }
        result = self._post("/deviceProfile", payload)
        log.info("  Created Device Profile '%s' (id=%s)", name, result["id"]["id"])
        return result

    # ------------------------------------------------------------------
    # Assets
    # ------------------------------------------------------------------
    def find_asset(self, name: str) -> Optional[Dict]:
        """Search for an asset by exact name."""
        if self.dry_run:
            return None
        try:
            data = self._get("/tenant/assets", params={"assetName": name})
            if data and data.get("name") == name:
                return data
        except requests.exceptions.HTTPError:
            pass
        return None

    def create_asset(self, name: str, asset_profile_id: str,
                     label: str = "", description: str = "") -> Dict:
        """Create or return existing asset."""
        existing = self.find_asset(name)
        if existing:
            log.info("  Asset '%s' already exists (id=%s)", name, existing["id"]["id"])
            return existing

        if self.dry_run:
            log.info("[DRY-RUN] Would create Asset: %s", name)
            return {"id": {"id": f"dry-run-asset-{name}", "entityType": "ASSET"}, "name": name}

        payload = {
            "name": name,
            "label": label or name,
            "assetProfileId": {"id": asset_profile_id, "entityType": "ASSET_PROFILE"},
        }
        if description:
            payload["additionalInfo"] = {"description": description}

        result = self._post("/asset", payload)
        log.info("  Created Asset '%s' (id=%s)", name, result["id"]["id"])
        return result

    # ------------------------------------------------------------------
    # Devices
    # ------------------------------------------------------------------
    def find_device(self, name: str) -> Optional[Dict]:
        """Search for a device by exact name."""
        if self.dry_run:
            return None
        try:
            data = self._get("/tenant/devices", params={"deviceName": name})
            if data and data.get("name") == name:
                return data
        except requests.exceptions.HTTPError:
            pass
        return None

    def create_device(self, name: str, device_profile_id: str,
                      label: str = "", description: str = "") -> Dict:
        """Create or return existing device."""
        existing = self.find_device(name)
        if existing:
            log.info("  Device '%s' already exists (id=%s)", name, existing["id"]["id"])
            return existing

        if self.dry_run:
            log.info("[DRY-RUN] Would create Device: %s", name)
            return {"id": {"id": f"dry-run-dev-{name}", "entityType": "DEVICE"}, "name": name}

        payload = {
            "name": name,
            "label": label or name,
            "deviceProfileId": {"id": device_profile_id, "entityType": "DEVICE_PROFILE"},
        }
        if description:
            payload["additionalInfo"] = {"description": description}

        result = self._post("/device", payload)
        log.info("  Created Device '%s' (id=%s)", name, result["id"]["id"])
        return result

    # ------------------------------------------------------------------
    # Relations
    # ------------------------------------------------------------------
    def create_relation(self, from_type: str, from_id: str,
                        to_type: str, to_id: str,
                        relation_type: str = "Contains") -> None:
        """Create a relation (FROM parent TO child). Idempotent -- TB ignores duplicates."""
        if self.dry_run:
            log.info("[DRY-RUN] Would create relation %s -> %s (%s)",
                     from_id[:8], to_id[:8], relation_type)
            return

        payload = {
            "from": {"id": from_id, "entityType": from_type},
            "to": {"id": to_id, "entityType": to_type},
            "type": relation_type,
            "typeGroup": "COMMON",
        }
        self._post("/relation", payload)
        log.info("    Relation %s [%s] -> %s [%s] (%s)",
                 from_type, from_id[:8], to_type, to_id[:8], relation_type)

    # ------------------------------------------------------------------
    # Attributes
    # ------------------------------------------------------------------
    def set_server_attributes(self, entity_type: str, entity_id: str,
                              attributes: Dict) -> None:
        """Set server-side attributes on an entity."""
        if self.dry_run:
            log.info("[DRY-RUN] Would set %d attributes on %s %s",
                     len(attributes), entity_type, entity_id[:8])
            return

        path = f"/plugins/telemetry/{entity_type}/{entity_id}/attributes/SERVER_SCOPE"
        self._post(path, attributes)
        log.info("    Set %d server attributes on %s %s",
                 len(attributes), entity_type, entity_id[:8])


# ===================================================================
# Alarm Rule Builder
# ===================================================================
def _build_alarm_rule(alarm_id: str, alarm_type: str, severity: str,
                      create_op: str, clear_op: str,
                      attribute: str) -> Dict:
    """Build a single alarm rule structure for the Device Profile."""
    return {
        "id": alarm_id,
        "alarmType": alarm_type,
        "createRules": {
            severity: {
                "condition": {
                    "condition": [
                        {
                            "key": {"type": "TIME_SERIES", "key": "PV"},
                            "valueType": "NUMERIC",
                            "value": None,
                            "predicate": {
                                "type": "NUMERIC",
                                "operation": create_op,
                                "value": {
                                    "defaultValue": 0,
                                    "userValue": None,
                                    "dynamicValue": {
                                        "sourceType": "CURRENT_DEVICE",
                                        "sourceAttribute": attribute,
                                        "inherit": True,
                                    },
                                },
                            },
                        }
                    ],
                    "spec": {"type": "SIMPLE"},
                },
                "schedule": None,
                "alarmDetails": None,
                "dashboardId": None,
            }
        },
        "clearRule": {
            "condition": {
                "condition": [
                    {
                        "key": {"type": "TIME_SERIES", "key": "PV"},
                        "valueType": "NUMERIC",
                        "value": None,
                        "predicate": {
                            "type": "NUMERIC",
                            "operation": clear_op,
                            "value": {
                                "defaultValue": 0,
                                "userValue": None,
                                "dynamicValue": {
                                    "sourceType": "CURRENT_DEVICE",
                                    "sourceAttribute": attribute,
                                    "inherit": True,
                                },
                            },
                        },
                    }
                ],
                "spec": {"type": "SIMPLE"},
            },
            "schedule": None,
            "alarmDetails": None,
            "dashboardId": None,
        },
        "propagate": True,
        "propagateToOwner": False,
        "propagateToOwnerHierarchy": False,
        "propagateToTenant": False,
        "propagateRelationTypes": None,
    }


def build_analog_alarms() -> List[Dict]:
    """Build the four HH/H/L/LL alarm rules for AnalogInput profile."""
    return [
        _build_alarm_rule(
            alarm_id="analog-hh-001",
            alarm_type="HIGH_HIGH",
            severity="CRITICAL",
            create_op="GREATER",
            clear_op="LESS",
            attribute="ss_alarmHH",
        ),
        _build_alarm_rule(
            alarm_id="analog-h-001",
            alarm_type="HIGH",
            severity="MAJOR",
            create_op="GREATER",
            clear_op="LESS",
            attribute="ss_alarmH",
        ),
        _build_alarm_rule(
            alarm_id="analog-l-001",
            alarm_type="LOW",
            severity="MAJOR",
            create_op="LESS",
            clear_op="GREATER",
            attribute="ss_alarmL",
        ),
        _build_alarm_rule(
            alarm_id="analog-ll-001",
            alarm_type="LOW_LOW",
            severity="CRITICAL",
            create_op="LESS",
            clear_op="GREATER",
            attribute="ss_alarmLL",
        ),
    ]


# ===================================================================
# Demo Data
# ===================================================================

# --- Asset Profile Definitions ---
ASSET_PROFILES = [
    {"name": "Enterprise", "description": "Top-level enterprise grouping"},
    {"name": "Plant", "description": "Physical plant / facility"},
    {"name": "Area", "description": "Functional area within a plant"},
    {"name": "Unit", "description": "Process unit / well / compressor train"},
    {"name": "Equipment", "description": "Individual equipment (pump, motor, etc.)"},
]

# --- Device Profile Definitions ---
DEVICE_PROFILES = [
    {"name": "AnalogInput", "description": "Analog input tag (4-20 mA, 0-10 V, etc.)", "alarms": "analog"},
    {"name": "DigitalInput", "description": "Digital input tag (on/off, run/stop)", "alarms": None},
    {"name": "Calculated", "description": "Calculated / derived tag", "alarms": None},
    {"name": "Totalizer", "description": "Totalizer / accumulator tag", "alarms": None},
]

# --- Demo Asset Hierarchy ---
# Each tuple: (name, profile, parent_name_or_None, description)
DEMO_ASSETS = [
    # Enterprise
    ("ACME Oil & Gas", "Enterprise", None, "Demo enterprise"),
    # Plant
    ("Planta Norte", "Plant", "ACME Oil & Gas", "Northern production facility"),
    # Areas
    ("Produccion ESP", "Area", "Planta Norte", "ESP production area"),
    ("Compresion", "Area", "Planta Norte", "Gas compression area"),
    # Units
    ("Pozo ESP-001", "Unit", "Produccion ESP", "ESP well 001"),
    ("Pozo ESP-002", "Unit", "Produccion ESP", "ESP well 002"),
    ("Compresor C-001", "Unit", "Compresion", "Compressor train C-001"),
    ("Compresor C-002", "Unit", "Compresion", "Compressor train C-002"),
    # Equipment
    ("Bomba-ESP-001", "Equipment", "Pozo ESP-001", "ESP pump 001"),
    ("Motor-ESP-001", "Equipment", "Pozo ESP-001", "ESP motor 001"),
    ("Bomba-ESP-002", "Equipment", "Pozo ESP-002", "ESP pump 002"),
    ("Motor-ESP-002", "Equipment", "Pozo ESP-002", "ESP motor 002"),
    ("Compresor-C001A", "Equipment", "Compresor C-001", "Compressor C-001 stage A"),
    ("Motor-C001", "Equipment", "Compresor C-001", "Compressor C-001 motor"),
    ("Compresor-C002A", "Equipment", "Compresor C-002", "Compressor C-002 stage A"),
    ("Motor-C002", "Equipment", "Compresor C-002", "Compressor C-002 motor"),
]

# --- Demo Tags (Devices) ---
# Each dict: tagName, deviceProfile, parentEquipment, and all attribute fields
DEMO_TAGS = [
    # === Pozo ESP-001 / Bomba-ESP-001 ===
    {
        "tagName": "ESP001_PIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Bomba-ESP-001",
        "ss_description": "ESP-001 Intake Pressure",
        "ss_engUnits": "PSI",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 5000,
        "ss_setpoint": 2500,
        "ss_deadband": 5,
        "ss_alarmHH": 4500,
        "ss_alarmH": 4000,
        "ss_alarmL": 500,
        "ss_alarmLL": 200,
        "ss_scanRate": "1s",
        "ss_instrumentType": "PIT",
    },
    {
        "tagName": "ESP001_FIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Bomba-ESP-001",
        "ss_description": "ESP-001 Flow Rate",
        "ss_engUnits": "BPD",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 10000,
        "ss_setpoint": 5000,
        "ss_deadband": 10,
        "ss_alarmHH": 9000,
        "ss_alarmH": 8000,
        "ss_alarmL": 1000,
        "ss_alarmLL": 500,
        "ss_scanRate": "1s",
        "ss_instrumentType": "FIT",
    },
    {
        "tagName": "ESP001_VIB_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Bomba-ESP-001",
        "ss_description": "ESP-001 Pump Vibration",
        "ss_engUnits": "g",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 10,
        "ss_setpoint": 0.5,
        "ss_deadband": 0.1,
        "ss_alarmHH": 8,
        "ss_alarmH": 5,
        "ss_alarmL": 0,
        "ss_alarmLL": 0,
        "ss_scanRate": "1s",
        "ss_instrumentType": "VIB",
    },
    # === Pozo ESP-001 / Motor-ESP-001 ===
    {
        "tagName": "ESP001_TIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Motor-ESP-001",
        "ss_description": "ESP-001 Motor Temperature",
        "ss_engUnits": "degF",
        "ss_dataType": "AI",
        "ss_rangeLow": 100,
        "ss_rangeHigh": 400,
        "ss_setpoint": 250,
        "ss_deadband": 1,
        "ss_alarmHH": 380,
        "ss_alarmH": 350,
        "ss_alarmL": 120,
        "ss_alarmLL": 110,
        "ss_scanRate": "1s",
        "ss_instrumentType": "TIT",
    },
    {
        "tagName": "ESP001_AMP_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Motor-ESP-001",
        "ss_description": "ESP-001 Motor Current",
        "ss_engUnits": "A",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 200,
        "ss_setpoint": 100,
        "ss_deadband": 0.5,
        "ss_alarmHH": 180,
        "ss_alarmH": 150,
        "ss_alarmL": 20,
        "ss_alarmLL": 10,
        "ss_scanRate": "1s",
        "ss_instrumentType": "AMP",
    },
    {
        "tagName": "ESP001_RUN_001",
        "deviceProfile": "DigitalInput",
        "parentEquipment": "Motor-ESP-001",
        "ss_description": "ESP-001 Motor Run Status",
        "ss_engUnits": "",
        "ss_dataType": "DI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 1,
        "ss_setpoint": 1,
        "ss_deadband": 0,
        "ss_scanRate": "1s",
        "ss_instrumentType": "XS",
    },
    # === Pozo ESP-002 / Bomba-ESP-002 ===
    {
        "tagName": "ESP002_PIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Bomba-ESP-002",
        "ss_description": "ESP-002 Intake Pressure",
        "ss_engUnits": "PSI",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 5000,
        "ss_setpoint": 2500,
        "ss_deadband": 5,
        "ss_alarmHH": 4500,
        "ss_alarmH": 4000,
        "ss_alarmL": 500,
        "ss_alarmLL": 200,
        "ss_scanRate": "1s",
        "ss_instrumentType": "PIT",
    },
    {
        "tagName": "ESP002_FIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Bomba-ESP-002",
        "ss_description": "ESP-002 Flow Rate",
        "ss_engUnits": "BPD",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 10000,
        "ss_setpoint": 5000,
        "ss_deadband": 10,
        "ss_alarmHH": 9000,
        "ss_alarmH": 8000,
        "ss_alarmL": 1000,
        "ss_alarmLL": 500,
        "ss_scanRate": "1s",
        "ss_instrumentType": "FIT",
    },
    {
        "tagName": "ESP002_VIB_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Bomba-ESP-002",
        "ss_description": "ESP-002 Pump Vibration",
        "ss_engUnits": "g",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 10,
        "ss_setpoint": 0.5,
        "ss_deadband": 0.1,
        "ss_alarmHH": 8,
        "ss_alarmH": 5,
        "ss_alarmL": 0,
        "ss_alarmLL": 0,
        "ss_scanRate": "1s",
        "ss_instrumentType": "VIB",
    },
    # === Pozo ESP-002 / Motor-ESP-002 ===
    {
        "tagName": "ESP002_TIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Motor-ESP-002",
        "ss_description": "ESP-002 Motor Temperature",
        "ss_engUnits": "degF",
        "ss_dataType": "AI",
        "ss_rangeLow": 100,
        "ss_rangeHigh": 400,
        "ss_setpoint": 250,
        "ss_deadband": 1,
        "ss_alarmHH": 380,
        "ss_alarmH": 350,
        "ss_alarmL": 120,
        "ss_alarmLL": 110,
        "ss_scanRate": "1s",
        "ss_instrumentType": "TIT",
    },
    {
        "tagName": "ESP002_AMP_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Motor-ESP-002",
        "ss_description": "ESP-002 Motor Current",
        "ss_engUnits": "A",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 200,
        "ss_setpoint": 100,
        "ss_deadband": 0.5,
        "ss_alarmHH": 180,
        "ss_alarmH": 150,
        "ss_alarmL": 20,
        "ss_alarmLL": 10,
        "ss_scanRate": "1s",
        "ss_instrumentType": "AMP",
    },
    {
        "tagName": "ESP002_RUN_001",
        "deviceProfile": "DigitalInput",
        "parentEquipment": "Motor-ESP-002",
        "ss_description": "ESP-002 Motor Run Status",
        "ss_engUnits": "",
        "ss_dataType": "DI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 1,
        "ss_setpoint": 1,
        "ss_deadband": 0,
        "ss_scanRate": "1s",
        "ss_instrumentType": "XS",
    },
    # === Compresor C-001 / Compresor-C001A ===
    {
        "tagName": "C001_PIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Compresor-C001A",
        "ss_description": "C-001 Suction Pressure",
        "ss_engUnits": "PSI",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 500,
        "ss_setpoint": 200,
        "ss_deadband": 2,
        "ss_alarmHH": 450,
        "ss_alarmH": 400,
        "ss_alarmL": 50,
        "ss_alarmLL": 20,
        "ss_scanRate": "1s",
        "ss_instrumentType": "PIT",
    },
    {
        "tagName": "C001_PIT_002",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Compresor-C001A",
        "ss_description": "C-001 Discharge Pressure",
        "ss_engUnits": "PSI",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 3000,
        "ss_setpoint": 1500,
        "ss_deadband": 5,
        "ss_alarmHH": 2700,
        "ss_alarmH": 2400,
        "ss_alarmL": 300,
        "ss_alarmLL": 150,
        "ss_scanRate": "1s",
        "ss_instrumentType": "PIT",
    },
    {
        "tagName": "C001_TIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Compresor-C001A",
        "ss_description": "C-001 Discharge Temperature",
        "ss_engUnits": "degF",
        "ss_dataType": "AI",
        "ss_rangeLow": 100,
        "ss_rangeHigh": 600,
        "ss_setpoint": 350,
        "ss_deadband": 2,
        "ss_alarmHH": 550,
        "ss_alarmH": 500,
        "ss_alarmL": 120,
        "ss_alarmLL": 110,
        "ss_scanRate": "1s",
        "ss_instrumentType": "TIT",
    },
    # === Compresor C-001 / Motor-C001 ===
    {
        "tagName": "C001_AMP_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Motor-C001",
        "ss_description": "C-001 Motor Current",
        "ss_engUnits": "A",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 500,
        "ss_setpoint": 250,
        "ss_deadband": 1,
        "ss_alarmHH": 450,
        "ss_alarmH": 400,
        "ss_alarmL": 50,
        "ss_alarmLL": 20,
        "ss_scanRate": "1s",
        "ss_instrumentType": "AMP",
    },
    {
        "tagName": "C001_VIB_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Motor-C001",
        "ss_description": "C-001 Motor Vibration",
        "ss_engUnits": "g",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 10,
        "ss_setpoint": 0.5,
        "ss_deadband": 0.1,
        "ss_alarmHH": 8,
        "ss_alarmH": 5,
        "ss_alarmL": 0,
        "ss_alarmLL": 0,
        "ss_scanRate": "1s",
        "ss_instrumentType": "VIB",
    },
    {
        "tagName": "C001_RUN_001",
        "deviceProfile": "DigitalInput",
        "parentEquipment": "Motor-C001",
        "ss_description": "C-001 Motor Run Status",
        "ss_engUnits": "",
        "ss_dataType": "DI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 1,
        "ss_setpoint": 1,
        "ss_deadband": 0,
        "ss_scanRate": "1s",
        "ss_instrumentType": "XS",
    },
    # === Compresor C-002 / Compresor-C002A ===
    {
        "tagName": "C002_PIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Compresor-C002A",
        "ss_description": "C-002 Suction Pressure",
        "ss_engUnits": "PSI",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 500,
        "ss_setpoint": 200,
        "ss_deadband": 2,
        "ss_alarmHH": 450,
        "ss_alarmH": 400,
        "ss_alarmL": 50,
        "ss_alarmLL": 20,
        "ss_scanRate": "1s",
        "ss_instrumentType": "PIT",
    },
    {
        "tagName": "C002_PIT_002",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Compresor-C002A",
        "ss_description": "C-002 Discharge Pressure",
        "ss_engUnits": "PSI",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 3000,
        "ss_setpoint": 1500,
        "ss_deadband": 5,
        "ss_alarmHH": 2700,
        "ss_alarmH": 2400,
        "ss_alarmL": 300,
        "ss_alarmLL": 150,
        "ss_scanRate": "1s",
        "ss_instrumentType": "PIT",
    },
    {
        "tagName": "C002_TIT_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Compresor-C002A",
        "ss_description": "C-002 Discharge Temperature",
        "ss_engUnits": "degF",
        "ss_dataType": "AI",
        "ss_rangeLow": 100,
        "ss_rangeHigh": 600,
        "ss_setpoint": 350,
        "ss_deadband": 2,
        "ss_alarmHH": 550,
        "ss_alarmH": 500,
        "ss_alarmL": 120,
        "ss_alarmLL": 110,
        "ss_scanRate": "1s",
        "ss_instrumentType": "TIT",
    },
    # === Compresor C-002 / Motor-C002 ===
    {
        "tagName": "C002_AMP_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Motor-C002",
        "ss_description": "C-002 Motor Current",
        "ss_engUnits": "A",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 500,
        "ss_setpoint": 250,
        "ss_deadband": 1,
        "ss_alarmHH": 450,
        "ss_alarmH": 400,
        "ss_alarmL": 50,
        "ss_alarmLL": 20,
        "ss_scanRate": "1s",
        "ss_instrumentType": "AMP",
    },
    {
        "tagName": "C002_VIB_001",
        "deviceProfile": "AnalogInput",
        "parentEquipment": "Motor-C002",
        "ss_description": "C-002 Motor Vibration",
        "ss_engUnits": "g",
        "ss_dataType": "AI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 10,
        "ss_setpoint": 0.5,
        "ss_deadband": 0.1,
        "ss_alarmHH": 8,
        "ss_alarmH": 5,
        "ss_alarmL": 0,
        "ss_alarmLL": 0,
        "ss_scanRate": "1s",
        "ss_instrumentType": "VIB",
    },
    {
        "tagName": "C002_RUN_001",
        "deviceProfile": "DigitalInput",
        "parentEquipment": "Motor-C002",
        "ss_description": "C-002 Motor Run Status",
        "ss_engUnits": "",
        "ss_dataType": "DI",
        "ss_rangeLow": 0,
        "ss_rangeHigh": 1,
        "ss_setpoint": 1,
        "ss_deadband": 0,
        "ss_scanRate": "1s",
        "ss_instrumentType": "XS",
    },
]


# ===================================================================
# Excel Loader
# ===================================================================
def load_from_excel(filepath: str) -> Tuple[List[Tuple], List[Dict]]:
    """
    Load asset hierarchy and tag definitions from an Excel file.

    Returns:
        (assets_list, tags_list) in the same format as DEMO_ASSETS and DEMO_TAGS.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        log.error("openpyxl is required for Excel import.  pip install openpyxl")
        sys.exit(1)

    wb = load_workbook(filepath, read_only=True, data_only=True)

    # --- Sheet "Assets" ---
    if "Assets" not in wb.sheetnames:
        log.error("Excel file must contain a sheet named 'Assets'")
        sys.exit(1)

    ws_assets = wb["Assets"]
    rows = list(ws_assets.iter_rows(min_row=1, values_only=True))
    header = [str(c).strip() if c else "" for c in rows[0]]
    expected_asset_cols = {"Nivel", "Nombre", "Perfil", "Padre", "Descripcion"}
    if not expected_asset_cols.issubset(set(header)):
        log.error("Assets sheet must have columns: %s  (found: %s)",
                  expected_asset_cols, header)
        sys.exit(1)

    idx = {name: i for i, name in enumerate(header)}
    assets_list: List[Tuple] = []
    for row in rows[1:]:
        if not row or not row[idx["Nombre"]]:
            continue
        name = str(row[idx["Nombre"]]).strip()
        profile = str(row[idx["Perfil"]]).strip()
        parent = str(row[idx["Padre"]]).strip() if row[idx["Padre"]] else None
        desc = str(row[idx["Descripcion"]]).strip() if row[idx["Descripcion"]] else ""
        assets_list.append((name, profile, parent, desc))

    # --- Sheet "Tags" ---
    if "Tags" not in wb.sheetnames:
        log.error("Excel file must contain a sheet named 'Tags'")
        sys.exit(1)

    ws_tags = wb["Tags"]
    trows = list(ws_tags.iter_rows(min_row=1, values_only=True))
    theader = [str(c).strip() if c else "" for c in trows[0]]
    expected_tag_cols = {"TagName", "DeviceProfile", "EquipoPadre", "EngUnits",
                         "RangeLow", "RangeHigh", "Deadband", "Setpoint",
                         "ScanRate", "InstrumentType", "Description"}
    if not expected_tag_cols.issubset(set(theader)):
        log.error("Tags sheet must have columns: %s  (found: %s)",
                  expected_tag_cols, theader)
        sys.exit(1)

    tidx = {name: i for i, name in enumerate(theader)}
    tags_list: List[Dict] = []
    for row in trows[1:]:
        if not row or not row[tidx["TagName"]]:
            continue
        tag_name = str(row[tidx["TagName"]]).strip()
        dev_profile = str(row[tidx["DeviceProfile"]]).strip()
        parent_eq = str(row[tidx["EquipoPadre"]]).strip()
        eng_units = str(row[tidx["EngUnits"]]).strip() if row[tidx["EngUnits"]] else ""
        range_low = float(row[tidx["RangeLow"]]) if row[tidx["RangeLow"]] is not None else 0
        range_high = float(row[tidx["RangeHigh"]]) if row[tidx["RangeHigh"]] is not None else 100
        deadband = float(row[tidx["Deadband"]]) if row[tidx["Deadband"]] is not None else 0
        setpoint = float(row[tidx["Setpoint"]]) if row[tidx["Setpoint"]] is not None else 0
        scan_rate = str(row[tidx["ScanRate"]]).strip() if row[tidx["ScanRate"]] else "1s"
        instr_type = str(row[tidx["InstrumentType"]]).strip() if row[tidx["InstrumentType"]] else ""
        description = str(row[tidx["Description"]]).strip() if row[tidx["Description"]] else ""

        tag: Dict[str, Any] = {
            "tagName": tag_name,
            "deviceProfile": dev_profile,
            "parentEquipment": parent_eq,
            "ss_description": description,
            "ss_engUnits": eng_units,
            "ss_dataType": "AI" if dev_profile == "AnalogInput" else "DI",
            "ss_rangeLow": range_low,
            "ss_rangeHigh": range_high,
            "ss_setpoint": setpoint,
            "ss_deadband": deadband,
            "ss_scanRate": scan_rate,
            "ss_instrumentType": instr_type,
        }

        # Alarm columns are optional (only for AnalogInput)
        if dev_profile == "AnalogInput":
            alarm_hh = row[tidx.get("AlarmHH", -1)] if "AlarmHH" in tidx and row[tidx["AlarmHH"]] is not None else 0
            alarm_h = row[tidx.get("AlarmH", -1)] if "AlarmH" in tidx and row[tidx["AlarmH"]] is not None else 0
            alarm_l = row[tidx.get("AlarmL", -1)] if "AlarmL" in tidx and row[tidx["AlarmL"]] is not None else 0
            alarm_ll = row[tidx.get("AlarmLL", -1)] if "AlarmLL" in tidx and row[tidx["AlarmLL"]] is not None else 0
            tag["ss_alarmHH"] = float(alarm_hh)
            tag["ss_alarmH"] = float(alarm_h)
            tag["ss_alarmL"] = float(alarm_l)
            tag["ss_alarmLL"] = float(alarm_ll)

        tags_list.append(tag)

    wb.close()
    log.info("Loaded %d assets and %d tags from %s", len(assets_list), len(tags_list), filepath)
    return assets_list, tags_list


# ===================================================================
# Sample Excel Generator
# ===================================================================
def generate_sample_excel(output_path: str) -> None:
    """Generate a sample_plant.xlsx from the built-in demo data."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        log.error("openpyxl is required.  pip install openpyxl")
        sys.exit(1)

    wb = Workbook()

    # --- Assets sheet ---
    ws_assets = wb.active
    ws_assets.title = "Assets"
    asset_headers = ["Nivel", "Nombre", "Perfil", "Padre", "Descripcion"]
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, color="FFFFFF")

    for col, hdr in enumerate(asset_headers, 1):
        cell = ws_assets.cell(row=1, column=col, value=hdr)
        cell.font = header_font_white
        cell.fill = header_fill

    # Map profile names to level numbers
    level_map = {"Enterprise": 1, "Plant": 2, "Area": 3, "Unit": 4, "Equipment": 5}
    for row_idx, (name, profile, parent, desc) in enumerate(DEMO_ASSETS, 2):
        ws_assets.cell(row=row_idx, column=1, value=level_map.get(profile, 0))
        ws_assets.cell(row=row_idx, column=2, value=name)
        ws_assets.cell(row=row_idx, column=3, value=profile)
        ws_assets.cell(row=row_idx, column=4, value=parent or "")
        ws_assets.cell(row=row_idx, column=5, value=desc)

    # Auto-width
    for col in ws_assets.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws_assets.column_dimensions[col_letter].width = max_len + 4

    # --- Tags sheet ---
    ws_tags = wb.create_sheet("Tags")
    tag_headers = ["TagName", "DeviceProfile", "EquipoPadre", "EngUnits",
                   "RangeLow", "RangeHigh", "Deadband", "Setpoint",
                   "AlarmHH", "AlarmH", "AlarmL", "AlarmLL",
                   "ScanRate", "InstrumentType", "Description"]
    for col, hdr in enumerate(tag_headers, 1):
        cell = ws_tags.cell(row=1, column=col, value=hdr)
        cell.font = header_font_white
        cell.fill = header_fill

    for row_idx, tag in enumerate(DEMO_TAGS, 2):
        ws_tags.cell(row=row_idx, column=1, value=tag["tagName"])
        ws_tags.cell(row=row_idx, column=2, value=tag["deviceProfile"])
        ws_tags.cell(row=row_idx, column=3, value=tag["parentEquipment"])
        ws_tags.cell(row=row_idx, column=4, value=tag["ss_engUnits"])
        ws_tags.cell(row=row_idx, column=5, value=tag["ss_rangeLow"])
        ws_tags.cell(row=row_idx, column=6, value=tag["ss_rangeHigh"])
        ws_tags.cell(row=row_idx, column=7, value=tag["ss_deadband"])
        ws_tags.cell(row=row_idx, column=8, value=tag["ss_setpoint"])
        ws_tags.cell(row=row_idx, column=9, value=tag.get("ss_alarmHH", ""))
        ws_tags.cell(row=row_idx, column=10, value=tag.get("ss_alarmH", ""))
        ws_tags.cell(row=row_idx, column=11, value=tag.get("ss_alarmL", ""))
        ws_tags.cell(row=row_idx, column=12, value=tag.get("ss_alarmLL", ""))
        ws_tags.cell(row=row_idx, column=13, value=tag["ss_scanRate"])
        ws_tags.cell(row=row_idx, column=14, value=tag["ss_instrumentType"])
        ws_tags.cell(row=row_idx, column=15, value=tag["ss_description"])

    # Auto-width
    for col in ws_tags.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws_tags.column_dimensions[col_letter].width = max_len + 4

    wb.save(output_path)
    log.info("Sample Excel file written to: %s", output_path)


# ===================================================================
# Build attributes dict from tag data
# ===================================================================
def build_attributes(tag: Dict) -> Dict:
    """Extract the ss_* keys from a tag dict for server-side attributes."""
    attrs = {}
    for key, val in tag.items():
        if key.startswith("ss_"):
            attrs[key] = val
    return attrs


# ===================================================================
# Main Provisioning Logic
# ===================================================================
def provision(client: TBClient, assets_data: List[Tuple], tags_data: List[Dict]) -> None:
    """
    Run the full provisioning pipeline:
      1. Asset Profiles
      2. Device Profiles (with alarm rules)
      3. Assets (hierarchy)
      4. Asset relations (Contains)
      5. Devices (tags)
      6. Device-to-Equipment relations (Contains)
      7. Server-side attributes on devices
    """

    # ------------------------------------------------------------------
    # Step 1: Asset Profiles
    # ------------------------------------------------------------------
    log.info("=" * 60)
    log.info("STEP 1: Asset Profiles")
    log.info("=" * 60)
    asset_profile_ids: Dict[str, str] = {}
    for ap in ASSET_PROFILES:
        result = client.create_asset_profile(ap["name"], ap["description"])
        asset_profile_ids[ap["name"]] = result["id"]["id"]
    log.info("Asset profiles ready: %d", len(asset_profile_ids))

    # ------------------------------------------------------------------
    # Step 2: Device Profiles
    # ------------------------------------------------------------------
    log.info("=" * 60)
    log.info("STEP 2: Device Profiles")
    log.info("=" * 60)
    device_profile_ids: Dict[str, str] = {}
    for dp in DEVICE_PROFILES:
        alarms = build_analog_alarms() if dp["alarms"] == "analog" else []
        result = client.create_device_profile(dp["name"], dp["description"], alarms)
        device_profile_ids[dp["name"]] = result["id"]["id"]
    log.info("Device profiles ready: %d", len(device_profile_ids))

    # ------------------------------------------------------------------
    # Step 3: Assets (hierarchy)
    # ------------------------------------------------------------------
    log.info("=" * 60)
    log.info("STEP 3: Assets (hierarchy)")
    log.info("=" * 60)
    asset_map: Dict[str, Dict] = {}  # name -> {id, entityType, ...}
    for name, profile, parent, desc in assets_data:
        ap_id = asset_profile_ids.get(profile)
        if not ap_id:
            log.error("Unknown asset profile '%s' for asset '%s'. Skipping.", profile, name)
            continue
        result = client.create_asset(name, ap_id, label=name, description=desc)
        asset_map[name] = result
    log.info("Assets ready: %d", len(asset_map))

    # ------------------------------------------------------------------
    # Step 4: Asset relations (Contains)
    # ------------------------------------------------------------------
    log.info("=" * 60)
    log.info("STEP 4: Asset Relations (Contains)")
    log.info("=" * 60)
    relations_created = 0
    for name, profile, parent, desc in assets_data:
        if parent and parent in asset_map and name in asset_map:
            parent_entity = asset_map[parent]
            child_entity = asset_map[name]
            client.create_relation(
                from_type="ASSET",
                from_id=parent_entity["id"]["id"],
                to_type="ASSET",
                to_id=child_entity["id"]["id"],
                relation_type="Contains",
            )
            relations_created += 1
    log.info("Asset relations created: %d", relations_created)

    # ------------------------------------------------------------------
    # Step 5: Devices (tags)
    # ------------------------------------------------------------------
    log.info("=" * 60)
    log.info("STEP 5: Devices (tags)")
    log.info("=" * 60)
    device_map: Dict[str, Dict] = {}  # tagName -> device entity
    for tag in tags_data:
        dp_id = device_profile_ids.get(tag["deviceProfile"])
        if not dp_id:
            log.error("Unknown device profile '%s' for tag '%s'. Skipping.",
                      tag["deviceProfile"], tag["tagName"])
            continue
        result = client.create_device(
            name=tag["tagName"],
            device_profile_id=dp_id,
            label=tag["tagName"],
            description=tag.get("ss_description", ""),
        )
        device_map[tag["tagName"]] = result
    log.info("Devices ready: %d", len(device_map))

    # ------------------------------------------------------------------
    # Step 6: Device-to-Equipment relations (Contains)
    # ------------------------------------------------------------------
    log.info("=" * 60)
    log.info("STEP 6: Equipment -> Device Relations (Contains)")
    log.info("=" * 60)
    dev_relations = 0
    for tag in tags_data:
        parent_name = tag["parentEquipment"]
        tag_name = tag["tagName"]
        if parent_name in asset_map and tag_name in device_map:
            parent_entity = asset_map[parent_name]
            device_entity = device_map[tag_name]
            client.create_relation(
                from_type="ASSET",
                from_id=parent_entity["id"]["id"],
                to_type="DEVICE",
                to_id=device_entity["id"]["id"],
                relation_type="Contains",
            )
            dev_relations += 1
        else:
            if parent_name not in asset_map:
                log.warning("Parent equipment '%s' not found for tag '%s'",
                            parent_name, tag_name)
    log.info("Device relations created: %d", dev_relations)

    # ------------------------------------------------------------------
    # Step 7: Server-side attributes
    # ------------------------------------------------------------------
    log.info("=" * 60)
    log.info("STEP 7: Server-side Attributes")
    log.info("=" * 60)
    attrs_set = 0
    for tag in tags_data:
        tag_name = tag["tagName"]
        if tag_name not in device_map:
            continue
        device_entity = device_map[tag_name]
        device_id = device_entity["id"]["id"]
        attributes = build_attributes(tag)
        client.set_server_attributes("DEVICE", device_id, attributes)
        attrs_set += 1
    log.info("Attributes set on %d devices", attrs_set)

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    log.info("=" * 60)
    log.info("PROVISIONING COMPLETE")
    log.info("=" * 60)
    log.info("  Asset Profiles : %d", len(asset_profile_ids))
    log.info("  Device Profiles: %d", len(device_profile_ids))
    log.info("  Assets         : %d", len(asset_map))
    log.info("  Asset Relations: %d", relations_created)
    log.info("  Devices (tags) : %d", len(device_map))
    log.info("  Dev Relations  : %d", dev_relations)
    log.info("  Attr Sets      : %d", attrs_set)


# ===================================================================
# CLI
# ===================================================================
def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Provisioner for ThingsBoard PE Industrial Historian",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python provisioner.py --demo
  python provisioner.py --file sample_plant.xlsx
  python provisioner.py --generate-sample
  python provisioner.py --demo --dry-run
        """,
    )

    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--demo", action="store_true",
                      help="Use built-in demo data (ACME Oil & Gas)")
    mode.add_argument("--file", type=str, metavar="FILE",
                      help="Read hierarchy and tags from an Excel (.xlsx) file")
    mode.add_argument("--generate-sample", action="store_true",
                      help="Generate sample_plant.xlsx in the scripts directory")

    parser.add_argument("--server", type=str, default="https://panel.atilax.io",
                        help="ThingsBoard server URL (default: https://panel.atilax.io)")
    parser.add_argument("--user", type=str, default="well@atilax.io",
                        help="ThingsBoard username (default: well@atilax.io)")
    parser.add_argument("--password", type=str, default="10203040",
                        help="ThingsBoard password (default: 10203040)")
    parser.add_argument("--dry-run", action="store_true",
                        help="Print what would be created without making API calls")

    return parser.parse_args()


def main() -> None:
    args = parse_args()

    # ------------------------------------------------------------------
    # Generate sample mode
    # ------------------------------------------------------------------
    if args.generate_sample:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        output_path = os.path.join(script_dir, "sample_plant.xlsx")
        generate_sample_excel(output_path)
        return

    # ------------------------------------------------------------------
    # Determine data source
    # ------------------------------------------------------------------
    if args.demo:
        log.info("Using built-in demo data (ACME Oil & Gas)")
        assets_data = DEMO_ASSETS
        tags_data = DEMO_TAGS
    elif args.file:
        if not os.path.isfile(args.file):
            log.error("File not found: %s", args.file)
            sys.exit(1)
        assets_data, tags_data = load_from_excel(args.file)
    else:
        log.error("Must specify --demo, --file, or --generate-sample")
        sys.exit(1)

    # ------------------------------------------------------------------
    # Validate data
    # ------------------------------------------------------------------
    log.info("Data loaded: %d assets, %d tags", len(assets_data), len(tags_data))

    # ------------------------------------------------------------------
    # Connect and provision
    # ------------------------------------------------------------------
    client = TBClient(
        base_url=args.server,
        username=args.user,
        password=args.password,
        dry_run=args.dry_run,
    )

    try:
        client.authenticate()
    except requests.exceptions.HTTPError as exc:
        log.error("Authentication failed: %s", exc)
        log.error("Response: %s", exc.response.text if exc.response else "N/A")
        sys.exit(1)
    except requests.exceptions.ConnectionError as exc:
        log.error("Cannot connect to %s: %s", args.server, exc)
        sys.exit(1)

    try:
        provision(client, assets_data, tags_data)
    except requests.exceptions.HTTPError as exc:
        log.error("API error: %s", exc)
        log.error("Response: %s", exc.response.text if exc.response else "N/A")
        sys.exit(1)
    except KeyboardInterrupt:
        log.warning("Interrupted by user")
        sys.exit(130)


if __name__ == "__main__":
    main()
